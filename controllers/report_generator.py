import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

class NumberedCanvas(canvas.Canvas):
    """Custom canvas to compute total page count and add headers/footers dynamically."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#475569")) # Slate grey
        
        # Header (Only on pages > 1 if we want, or all)
        self.drawString(36, 756, "Smart Faculty Invigilation Allotment System")
        self.drawRightString(doc_width - 36, 756, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        self.setStrokeColor(colors.HexColor("#CBD5E1")) # Light border
        self.setLineWidth(0.5)
        self.line(36, 748, doc_width - 36, 748)
        
        # Footer
        self.line(36, 45, doc_width - 36, 45)
        self.drawString(36, 30, "College Examination Cell - Confidential")
        self.drawRightString(doc_width - 36, 30, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()

# Global page size settings for PDF
doc_width, doc_height = letter

def generate_excel_report(data_rows, headers, report_title="Report"):
    """
    Generates a stylized Excel sheet.
    data_rows: list of lists
    headers: list of strings
    """
    wb = Workbook()
    ws = wb.active
    ws.title = report_title[:30] # Excel limit is 31 chars
    ws.views.sheetView[0].showGridLines = True
    
    # Fonts
    title_font = Font(name='Segoe UI', size=16, bold=True, color='1E293B')
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    body_font = Font(name='Segoe UI', size=10, color='334155')
    
    # Fills & Borders
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid') # Deep blue
    border_side = Side(border_style='thin', color='E2E8F0')
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    # Title Row
    ws.append([report_title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(row=1, column=1).font = title_font
    ws.cell(row=1, column=1).alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 40
    
    # Spacer row
    ws.append([])
    ws.row_dimensions[2].height = 15
    
    # Headers Row
    ws.append(headers)
    ws.row_dimensions[3].height = 25
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = cell_border
        
    # Data Rows
    for row_idx, row in enumerate(data_rows, 4):
        ws.append(row)
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = body_font
            cell.border = cell_border
            # Format alignment based on content type
            val = cell.value
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal='right', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Avoid measuring title row
            if cell.row == 1:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def generate_pdf_report(data_rows, headers, report_title="Report"):
    """
    Generates a beautifully designed PDF document.
    data_rows: list of lists
    headers: list of strings
    """
    buffer = io.BytesIO()
    
    # Page setup
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=72, # Room for custom header in NumberedCanvas
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0F172A'),
        alignment=0, # Left align
        spaceAfter=15
    )
    
    header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=12,
        textColor=colors.white,
        alignment=1 # Center
    )
    
    body_style = ParagraphStyle(
        'TableBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=11,
        textColor=colors.HexColor('#334155')
    )
    
    story = []
    
    # Title
    story.append(Paragraph(report_title, title_style))
    story.append(Spacer(1, 10))
    
    # Build Table
    table_data = []
    # Headers
    table_data.append([Paragraph(h, header_style) for h in headers])
    # Rows
    for row in data_rows:
        table_row = []
        for cell in row:
            val_str = str(cell) if cell is not None else ""
            table_row.append(Paragraph(val_str, body_style))
        table_data.append(table_row)
        
    # Table sizing
    # Calculate even column width based on margins
    usable_width = doc_width - doc.leftMargin - doc.rightMargin
    col_width = usable_width / len(headers)
    
    pdf_table = Table(table_data, colWidths=[col_width] * len(headers))
    
    # Table styling
    t_style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')), # Navy blue header
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')), # Light grey borders
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]), # Alternating background rows
        ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
        ('TOPPADDING', (0, 1), (-1, -1), 6),
    ])
    pdf_table.setStyle(t_style)
    
    story.append(pdf_table)
    
    # Build document using NumberedCanvas
    doc.build(story, canvasmaker=NumberedCanvas)
    
    buffer.seek(0)
    return buffer
